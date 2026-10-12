import os
import sys
import json
import io

# Ensure repository root is on sys.path regardless of working directory
_TEST_DIR = os.path.dirname(os.path.abspath(__file__))
_PARENT_DIR = os.path.abspath(os.path.join(_TEST_DIR, ".."))
for _p in [_TEST_DIR, _PARENT_DIR]:
    if _p not in sys.path:
        sys.path.insert(0, _p)

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl
from compile_attendance_excel import (
    grab,
    is_granted,
    normalize_json_record,
    normalize_text_record,
    add_record,
    autosize_and_style,
    lambda_handler,
    HEADERS
)

print("=" * 70)
print("RUNNING COMPREHENSIVE EXCEL COMPILER VERIFICATION")
print("=" * 70)

# 1. Test grab() helper
sample_text = (
    "Name: John Doe\n"
    "Time (UTC): 2026-10-12T12:00:00Z\n"
    "Typing Time (s): 3.45\n"
    "Platform: Windows\n"
    "Language: en-US\n"
    "Screen: 15.6 inches\n"
    "Timezone: Africa/Lusaka\n"
    "Connection: 4G\n"
    "IP: 197.234.12.1\n"
    "Camera Permission Status: granted\n"
    "Permission Status: granted\n"
    "Latitude: -15.4167\n"
    "Longitude: 28.2833\n"
    "Accuracy Meters: 10\n"
    "State: Lusaka Province\n"
    "City: Lusaka\n"
    "Street Name: Great East Rd\n"
    "Altitude: 1200\n"
    "Altitude Accuracy: 5\n"
    "Heading: 0\n"
    "Speed: 0\n"
    "Location Timestamp: 2026-10-12T12:00:00Z\n"
)

assert grab("Name", sample_text) == "John Doe", f"grab Name failed: {grab('Name', sample_text)}"
assert grab("Latitude", sample_text) == "-15.4167", f"grab Latitude failed: {grab('Latitude', sample_text)}"
assert grab("NonExistent", sample_text) == "", f"grab NonExistent failed: {grab('NonExistent', sample_text)}"
print("✓ grab() parsing verified")

# 2. Test is_granted()
assert is_granted("granted") is True
assert is_granted("Granted") is True
assert is_granted(" GRANTED ") is True
assert is_granted("denied") is False
assert is_granted("") is False
assert is_granted(None) is False
print("✓ is_granted() status validation verified")

# 3. Test normalize_json_record()
sample_json = {
    "name": "Jane Smith",
    "timestamp": "2026-10-12T12:05:00Z",
    "typingTimeSeconds": "2.10",
    "platform": "MacBook / macOS",
    "language": "en-GB",
    "screen": "13.3 inches",
    "timezone": "Europe/London",
    "network": "WiFi",
    "ip": "82.165.197.1",
    "cameraPermission": "granted",
    "locationPermission": "denied",
    "location": {
        "latitude": "",
        "longitude": "",
        "accuracy": "",
        "state": "",
        "city": "",
        "streetName": "",
        "altitude": "",
        "altitudeAccuracy": "",
        "heading": "",
        "speed": "",
        "timestamp": ""
    }
}
norm_json = normalize_json_record(sample_json)
assert norm_json["Name"] == "Jane Smith"
assert norm_json["CameraPermission"] == "granted"
assert norm_json["LocationPermission"] == "denied"
assert norm_json["Platform"] == "MacBook / macOS"

# Non-dict location fallback
malformed = {"name": "Broken User", "location": "not a dict"}
norm_malformed = normalize_json_record(malformed)
assert norm_malformed["Name"] == "Broken User"
assert norm_malformed["Latitude"] == ""
print("✓ normalize_json_record() schema mapping verified")

# 4. Test normalize_text_record()
norm_txt = normalize_text_record(sample_text)
assert norm_txt["Name"] == "John Doe"
assert norm_txt["City"] == "Lusaka"
assert norm_txt["CameraPermission"] == "granted"
assert norm_txt["LocationPermission"] == "granted"
print("✓ normalize_text_record() schema mapping verified")

# 5. Test add_record() deduplication
records = {}
add_record(records, norm_json)
assert len(records) == 1
# Duplicate record with same name and timestamp should overwrite, not duplicate
add_record(records, norm_json)
assert len(records) == 1
# Record with different timestamp should be added
diff_time = dict(norm_json)
diff_time["Time"] = "2026-10-12T13:00:00Z"
add_record(records, diff_time)
assert len(records) == 2
# Blank name record should be skipped
add_record(records, {"Name": "", "Time": "2026-10-12T14:00:00Z"})
assert len(records) == 2
print("✓ add_record() deduplication and key collision handling verified")

# 6. Test autosize_and_style()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test"
ws.append(["Col1", "Col2"])
ws.append(["Value 1", "Much Longer Value Two"])
autosize_and_style(ws)
assert ws["A1"].font.bold is True
assert ws.column_dimensions["A"].width >= 14
print("✓ autosize_and_style() formatting verified")

# 7. Test lambda_handler() end-to-end with mock S3
class MockBody:
    def __init__(self, data_str):
        self.data_str = data_str
    def read(self):
        return self.data_str.encode("utf-8")

class MockPaginator:
    def paginate(self, Bucket, Prefix):
        if Prefix == "records/":
            return [{
                "Contents": [
                    {"Key": "records/2026-10-12/jane_smith.json"}
                ]
            }]
        elif Prefix == "raw-text/":
            return [{
                "Contents": [
                    {"Key": "raw-text/2026-10-12/john_doe.txt"}
                ]
            }]
        return []

uploaded_s3_files = {}

class MockS3Client:
    def get_paginator(self, name):
        return MockPaginator()

    def get_object(self, Bucket, Key):
        if Key == "records/2026-10-12/jane_smith.json":
            return {"Body": MockBody(json.dumps(sample_json))}
        elif Key == "raw-text/2026-10-12/john_doe.txt":
            return {"Body": MockBody(sample_text)}
        raise KeyError(f"Key not found: {Key}")

    def put_object(self, Bucket, Key, Body, ContentType):
        uploaded_s3_files[Key] = {
            "bucket": Bucket,
            "bytes": Body.read() if hasattr(Body, "read") else Body,
            "content_type": ContentType
        }

# Patch boto3 s3 client in compile_attendance_excel
import compile_attendance_excel
old_s3 = compile_attendance_excel.s3
compile_attendance_excel.s3 = MockS3Client()

try:
    result = lambda_handler({}, None)
    assert result["status"] == "updated", f"Unexpected result status: {result}"
    assert result["totalAttendance"] == 2, f"Expected 2 records, got {result['totalAttendance']}"
    assert result["cameraPermissionGranted"] == 2, f"Expected 2 camera permissions granted, got {result['cameraPermissionGranted']}"
    assert result["locationPermissionGranted"] == 1, f"Expected 1 location permission granted, got {result['locationPermissionGranted']}"

    # Verify generated Excel content
    excel_key = compile_attendance_excel.OUTPUT_KEY
    assert excel_key in uploaded_s3_files, f"OUTPUT_KEY {excel_key} was not uploaded to S3"

    excel_bytes = uploaded_s3_files[excel_key]["bytes"]
    wb_loaded = openpyxl.load_workbook(io.BytesIO(excel_bytes))

    sheet_names = wb_loaded.sheetnames
    assert "Attendance" in sheet_names, f"Attendance sheet missing from {sheet_names}"
    assert "Summary" in sheet_names, f"Summary sheet missing from {sheet_names}"

    att_ws = wb_loaded["Attendance"]
    assert att_ws.max_row == 3  # 1 header + 2 records
    assert att_ws.cell(row=1, column=1).value == "S No"
    assert att_ws.cell(row=1, column=2).value == "Name"

    sum_ws = wb_loaded["Summary"]
    assert sum_ws.max_row == 6  # header + 5 summary metrics
    assert sum_ws.cell(row=2, column=2).value == 2  # Total Attendance Marked
    assert sum_ws.cell(row=3, column=2).value == 2  # Camera Permission Granted
    assert sum_ws.cell(row=4, column=2).value == 1  # Location Permission Granted

    print("✓ lambda_handler() end-to-end execution, Excel generation, and metrics verified")
finally:
    compile_attendance_excel.s3 = old_s3

print("=" * 70)
print("ALL EXCEL COMPILER TESTS PASSED SUCCESSFULLY!")
print("=" * 70)
